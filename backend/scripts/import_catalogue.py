"""Import an admin-reviewed CSV/XLSX catalogue into Supabase/PostgreSQL.

Usage: python scripts/import_catalogue.py path/to/catalogue.csv
All imported rows remain PENDING_ADMIN_REVIEW. Only a privileged admin should
change their verification_status to OFFICIAL_SOURCE after checking each URL.
"""
import csv
import json
import os
import sys
from datetime import date
from pathlib import Path

import psycopg

REQUIRED = {"name", "scheme_service_type", "government_level", "state_ut", "department", "category", "description", "official_source_url", "application_url", "last_verified_date", "verification_status", "is_active"}
LEVELS = {"CENTRAL", "STATE", "UT"}
TYPES = {"SCHEME", "SERVICE"}

def rows_from(path: Path):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f: return list(csv.DictReader(f))
    if path.suffix.lower() in {".xlsx", ".xls"}:
        from openpyxl import load_workbook
        sheet = load_workbook(path, read_only=True, data_only=True).active
        headers = [str(x.value or "").strip() for x in next(sheet.iter_rows(max_row=1))]
        return [dict(zip(headers, [("" if c.value is None else str(c.value)) for c in row])) for row in sheet.iter_rows(min_row=2)]
    raise ValueError("Use a .csv or .xlsx file")

def normalise(row, number):
    row = {str(k).strip(): (v or "").strip() for k, v in row.items()}
    missing = REQUIRED - row.keys()
    if missing: raise ValueError(f"row {number}: missing columns: {', '.join(sorted(missing))}")
    if not row["name"] or not row["official_source_url"]: raise ValueError(f"row {number}: name and official_source_url are required")
    if row["government_level"].upper() not in LEVELS: raise ValueError(f"row {number}: invalid government_level")
    if row["scheme_service_type"].upper() not in TYPES: raise ValueError(f"row {number}: invalid scheme_service_type")
    if row["government_level"].upper() != "CENTRAL" and not row["state_ut"]: raise ValueError(f"row {number}: state_ut is required for State/UT records")
    try: date.fromisoformat(row["last_verified_date"])
    except ValueError: raise ValueError(f"row {number}: last_verified_date must be YYYY-MM-DD")
    return row

def main():
    if len(sys.argv) != 2: raise SystemExit("Usage: python scripts/import_catalogue.py catalogue.csv|xlsx")
    database_url = os.environ.get("DATABASE_URL")
    if not database_url: raise SystemExit("DATABASE_URL must be set; do not put it in the import file.")
    records = [normalise(row, i) for i, row in enumerate(rows_from(Path(sys.argv[1])), start=2)]
    sql = """with department as (insert into departments (name,official_url) values (%(department)s,%(official_source_url)s) on conflict (name) do update set official_url=excluded.official_url returning id)
    insert into schemes (department_id,name,description,category,state,benefits,eligibility,official_url,application_url,information_status,last_verified_at,active,scheme_service_type,government_level,state_ut_code,verification_status,source_reference)
    values ((select id from department),%(name)s,%(description)s,%(category)s,nullif(%(state_ut)s,''),%(benefits)s::jsonb,%(eligibility)s::jsonb,%(official_source_url)s,nullif(%(application_url)s,''),%(verification_status)s,%(last_verified_date)s,%(is_active)s,%(scheme_service_type)s,%(government_level)s,(select code from jurisdictions where name=%(state_ut)s), 'PENDING_ADMIN_REVIEW',%(official_source_url)s) returning id"""
    with psycopg.connect(database_url) as conn:
        with conn.cursor() as cur:
            for row in records:
                row.update(benefits=json.dumps([x.strip() for x in row.get("benefits", "").split("|") if x.strip()]), eligibility=row.get("eligibility", "{}"), is_active=row["is_active"].lower() in {"true","1","yes"})
                cur.execute(sql, row)
                scheme_id = cur.fetchone()[0]
                for name in [x.strip() for x in row.get("required_documents", "").split("|") if x.strip()]:
                    cur.execute("insert into scheme_documents (scheme_id,name) values (%s,%s)", (scheme_id, name))
    print(f"Imported {len(records)} records as PENDING_ADMIN_REVIEW.")

if __name__ == "__main__": main()
